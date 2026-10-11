"""Step 3 - Run the SQL revenue queries, build the forecast, write the findings.

Outputs:
  outputs/sql_results/*.csv     one CSV per SQL query
  outputs/forecast.csv          next 3 months with low/high range
  outputs/forecast_backtest.csv every backtest prediction and its error
  outputs/revenue_report.xlsx   formatted Excel report
  outputs/insights.md           key findings
"""
import re
import sqlite3

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import forecast as fc
from config import CURRENCY, DB_PATH, EXCEL_REPORT, FORECAST_CSV, OUTPUT_DIR, QUERIES_SQL

PLAIN_METHOD = {
    "Naive (last month)": "the naive method (next month = last month)",
    "3-month moving average": "the 3-month moving average",
    "Linear trend": "the linear trend",
}

MONEY_COLS = {"revenue", "avg_order_value", "freight", "amount_paid", "revenue_jan_aug",
              "avg_item_price", "actual", "predicted", "forecast", "low", "high"}


def load_queries(path) -> dict[str, str]:
    parts = re.split(r"^-- name:\s*(\w+)\s*$", path.read_text(), flags=re.MULTILINE)
    return {name: sql.strip() for name, sql in zip(parts[1::2], parts[2::2])}


def run_queries() -> dict[str, pd.DataFrame]:
    with sqlite3.connect(DB_PATH) as conn:
        return {name: pd.read_sql_query(sql, conn) for name, sql in load_queries(QUERIES_SQL).items()}


def money(x: float) -> str:
    """R$13,181,027 -> 'R$13.18M'; R$987,765 -> 'R$987.8K'"""
    if abs(x) >= 1e6:
        return f"{CURRENCY}{x / 1e6:.2f}M"
    if abs(x) >= 1e3:
        return f"{CURRENCY}{x / 1e3:.1f}K"
    return f"{CURRENCY}{x:,.0f}"


def month_name(m: str) -> str:
    return pd.Period(m, freq="M").strftime("%b %Y")


def build_insights(r: dict[str, pd.DataFrame], f: dict) -> list[str]:
    o = r["overall_kpis"].iloc[0]
    mon = r["monthly_revenue"]
    yc = r["year_comparison"].set_index("year")
    cats = r["category_revenue"]
    par = r["product_pareto"].iloc[0]
    states = r["state_revenue"]
    days = r["top_days"]
    loyal = r["customer_loyalty"].set_index("customer_group")
    wd = r["weekday_revenue"]
    pay = r["payment_methods"].iloc[0]

    y17, y18 = yc.loc["2017", "revenue_jan_aug"], yc.loc["2018", "revenue_jan_aug"]
    m18 = mon[mon["month"] >= "2018-01"]
    peak = m18.loc[m18["revenue"].idxmax()]
    last = mon.iloc[-1]
    nov = mon[mon["month"] == "2017-11"].iloc[0]
    cats_80 = int((cats["cumulative_share_pct"] < 80).sum() + 1)
    top3_states = states.head(3)
    big_states = states[states["orders"] >= 1000]
    costly = big_states.loc[big_states["freight_pct_of_revenue"].idxmax()]
    repeat_customers = 100 - loyal.loc["1 order", "customers_pct"]
    repeat_revenue = 100 - loyal.loc["1 order", "revenue_pct"]
    best_day, worst_day = wd.loc[wd["share_pct"].idxmax()], wd.loc[wd["share_pct"].idxmin()]
    fcast = f["forecast"]
    trend = f["scores"].set_index("method").loc["Linear trend", "mape_pct"]

    lines = [
        f"**Overall:** {money(o.revenue)} revenue from {int(o.orders):,} delivered orders and "
        f"{int(o.customers):,} customers (Jan 2017 - Aug 2018). Average order value: {CURRENCY}{o.avg_order_value:.0f}.",

        f"**Growth then a plateau:** Jan-Aug 2018 revenue was {money(y18)}, **{(y18 / y17 - 1) * 100:.0f}% higher** "
        f"than Jan-Aug 2017 ({money(y17)}). But growth stopped in 2018: monthly revenue peaked at {money(peak.revenue)} in "
        f"{month_name(peak.month)} and was {money(last.revenue)} by {month_name(last.month)} "
        f"({(last.revenue / peak.revenue - 1) * 100:.0f}%).",

        f"**Black Friday:** {month_name(nov.month)} revenue jumped **{nov.mom_growth_pct:.0f}%** over the month before. "
        f"The single best day, {pd.Timestamp(days.iloc[0].order_date):%d %b %Y}, brought {money(days.iloc[0].revenue)} - "
        f"{days.iloc[0].revenue / days.iloc[1].revenue:.1f}x the next best day.",

        f"**Categories:** {len(cats)} categories sell, but {cats_80} of them make 80% of revenue. The top 3 "
        f"({', '.join(cats.head(3)['category'])}) bring {cats.head(3)['share_pct'].sum():.1f}%.",

        f"**80/20 rule:** {par.pct_of_products:.0f}% of products ({int(par.products_for_80pct_revenue):,} of "
        f"{int(par.products_sold):,}) generate 80% of revenue.",

        f"**Geography:** {', '.join(top3_states['state'])} bring {top3_states['share_pct'].sum():.0f}% of revenue "
        f"({top3_states.iloc[0]['state']} alone {top3_states.iloc[0]['share_pct']:.0f}%). Shipping is costliest in "
        f"{costly.state}, where freight equals {costly.freight_pct_of_revenue:.0f}% of the order value "
        f"(vs {o.freight_pct_of_revenue:.0f}% overall).",

        f"**Loyalty:** only **{repeat_customers:.1f}%** of customers ordered more than once, and they bring just "
        f"{repeat_revenue:.1f}% of revenue. Almost all revenue depends on finding new customers.",

        f"**Weekdays:** {best_day.day} is the best day ({best_day.share_pct:.1f}% of revenue); "
        f"{worst_day.day} is the weakest ({worst_day.share_pct:.1f}%). Weekends sell less than any weekday.",

        f"**Payments:** {pay.payment_type.lower()} pays for {pay.share_pct:.0f}% of the money, "
        f"spread over {pay.avg_installments:.1f} instalments on average.",

        f"**Forecast:** the backtest picked **{PLAIN_METHOD[f['best_method']]}** (average error {f['best_mape']:.1f}%); a straight-line "
        f"trend was off by {trend:.1f}% because growth stopped in 2018. Expected revenue is about "
        f"**{money(fcast.iloc[0].forecast)} a month** for {month_name(fcast.iloc[0].month)} - "
        f"{month_name(fcast.iloc[-1].month)}, with a range of {money(fcast.iloc[0].low)}-{money(fcast.iloc[0].high)} "
        f"next month widening to {money(fcast.iloc[-1].low)}-{money(fcast.iloc[-1].high)} three months out. "
        f"November will likely beat this because of Black Friday.",
    ]

    md = ["# Key Insights", "", *[f"{i}. {s}" for i, s in enumerate(lines, 1)], "",
          "_Revenue = item prices of delivered orders; freight is excluded. Amounts in Brazilian reais (R$)._", ""]
    (OUTPUT_DIR / "insights.md").write_text("\n".join(md), encoding="utf-8")
    return lines


def write_excel(sheets: dict[str, pd.DataFrame]):
    with pd.ExcelWriter(EXCEL_REPORT, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2A78D6")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(df.columns, 1):
                letter = get_column_letter(i)
                width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
                ws.column_dimensions[letter].width = min(width, 90)
                if col in MONEY_COLS:
                    for cell in ws[letter][1:]:
                        cell.number_format = '"R$"#,##0.00'


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    r = run_queries()
    f = fc.run(r["monthly_revenue"][["month", "revenue"]])

    sql_dir = OUTPUT_DIR / "sql_results"
    sql_dir.mkdir(exist_ok=True)
    for name, df in r.items():
        df.to_csv(sql_dir / f"{name}.csv", index=False)
    f["forecast"].to_csv(FORECAST_CSV, index=False)
    f["backtest"].to_csv(OUTPUT_DIR / "forecast_backtest.csv", index=False)
    f["scores"].to_csv(OUTPUT_DIR / "forecast_method_scores.csv", index=False)

    insights = build_insights(r, f)
    write_excel({
        "Summary": r["overall_kpis"],
        "Monthly Revenue": r["monthly_revenue"],
        "2017 vs 2018": r["year_comparison"],
        "Forecast": f["forecast"],
        "Forecast Methods": f["scores"],
        "Forecast Backtest": f["backtest"],
        "Categories": r["category_revenue"],
        "Product Pareto": r["product_pareto"],
        "States": r["state_revenue"],
        "Weekdays": r["weekday_revenue"],
        "Top Days": r["top_days"],
        "Customer Loyalty": r["customer_loyalty"],
        "Payments": r["payment_methods"],
        "Insights": pd.DataFrame({"insight": [s.replace("**", "") for s in insights]}),
    })

    print(f"[3/5] Ran {len(r)} SQL queries -> outputs/sql_results/, Excel report, insights")
    print(f"[4/5] Forecast: backtested {len(fc.METHODS)} methods, best = {f['best_method']} "
          f"(avg error {f['best_mape']:.1f}%) -> outputs/forecast.csv")
    print()
    for line in insights:
        print("      - " + line.replace("**", ""))
    print()


if __name__ == "__main__":
    main()
